"""
Gestionnaire d'export Excel et PDF.

Ce module permet d'exporter les tableaux de prédictions et les graphiques
en formats Excel et PDF pour partage et archivage.
"""

import pandas as pd
import matplotlib.pyplot as plt
import matplotlib.dates as mdates
from datetime import datetime
from pathlib import Path
from typing import Dict, Optional
import os


class ExportManager:
    """Gestionnaire d'export Excel et PDF"""
    
    def __init__(self, export_dir: str = "exports"):
        """
        Initialise le gestionnaire d'export
        
        Args:
            export_dir: Répertoire de destination des exports
        """
        self.export_dir = Path(export_dir)
        self.export_dir.mkdir(exist_ok=True)
    
    def export_to_excel(self, data: pd.DataFrame, filename: str = None, metadata: Dict = None) -> str:
        """
        Exporte les données en Excel
        
        Args:
            data: DataFrame à exporter
            filename: Nom du fichier (généré automatiquement si None)
            metadata: Métadonnées (k(A), ε, date génération)
            
        Returns:
            Chemin du fichier exporté
        """
        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment
        except ImportError:
            raise ImportError("openpyxl est requis pour l'export Excel. Installez-le avec: pip install openpyxl")
        
        # Générer le nom de fichier si non fourni
        if filename is None:
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            filename = f"predictions_formule_{timestamp}.xlsx"
        
        # Ajouter l'extension si nécessaire
        if not filename.endswith('.xlsx'):
            filename += '.xlsx'
        
        filepath = self.export_dir / filename
        
        # Créer le workbook
        with pd.ExcelWriter(filepath, engine='openpyxl') as writer:
            # Écrire les données
            data.to_excel(writer, sheet_name='Prédictions', index=False)
            
            # Obtenir le workbook et la feuille
            workbook = writer.book
            worksheet = writer.sheets['Prédictions']
            
            # Formater l'en-tête
            header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
            header_font = Font(color="FFFFFF", bold=True)
            
            for cell in worksheet[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal='center', vertical='center')
            
            # Ajuster la largeur des colonnes
            for column in worksheet.columns:
                max_length = 0
                column_letter = column[0].column_letter
                
                for cell in column:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                    except:
                        pass
                
                adjusted_width = min(max_length + 2, 50)
                worksheet.column_dimensions[column_letter].width = adjusted_width
            
            # Ajouter les métadonnées si fournies
            if metadata:
                metadata_sheet = workbook.create_sheet('Métadonnées')
                
                row = 1
                for key, value in metadata.items():
                    metadata_sheet.cell(row=row, column=1, value=key)
                    metadata_sheet.cell(row=row, column=2, value=str(value))
                    row += 1
                
                # Formater les métadonnées
                for cell in metadata_sheet['A']:
                    cell.font = Font(bold=True)
        
        return str(filepath)
    
    def export_to_pdf(self, data: pd.DataFrame, filename: str = None, metadata: Dict = None) -> str:
        """
        Exporte les données en PDF avec graphique
        
        Args:
            data: DataFrame à exporter
            filename: Nom du fichier (généré automatiquement si None)
            metadata: Métadonnées (k(A), ε, date génération)
            
        Returns:
            Chemin du fichier exporté
        """
        # Générer le nom de fichier si non fourni
        if filename is None:
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            filename = f"predictions_formule_{timestamp}.pdf"
        
        # Ajouter l'extension si nécessaire
        if not filename.endswith('.pdf'):
            filename += '.pdf'
        
        filepath = self.export_dir / filename
        
        # Créer le graphique
        chart_path = self.create_prediction_chart(data, metadata)
        
        # Pour l'instant, on sauvegarde juste le graphique en PDF
        # Une implémentation complète nécessiterait reportlab ou matplotlib.backends.backend_pdf
        
        # Sauvegarder le graphique comme PDF
        fig = plt.gcf()
        fig.savefig(filepath, format='pdf', bbox_inches='tight', dpi=300)
        plt.close(fig)
        
        return str(filepath)
    
    def create_prediction_chart(self, data: pd.DataFrame, metadata: Dict = None) -> str:
        """
        Crée un graphique des prédictions avec matplotlib
        
        Args:
            data: DataFrame avec les prédictions
            metadata: Métadonnées pour le titre
            
        Returns:
            Chemin du fichier image généré
        """
        # Créer la figure
        fig, ax = plt.subplots(figsize=(14, 8))
        
        # Extraire les données
        dates = pd.to_datetime(data['Date'])
        Q_central = data['Q centrale']
        Q_min = data['Q min (-8%)']
        Q_max = data['Q max (+19%)']
        
        # Tracer les courbes
        ax.plot(dates, Q_central, 'b-', linewidth=2, label='Q centrale')
        ax.fill_between(dates, Q_min, Q_max, alpha=0.3, color='blue', label='Intervalle -8%/+19%')
        
        # Si Q réelle existe, la tracer
        if 'Q réelle' in data.columns and data['Q réelle'].notna().any():
            Q_real = data['Q réelle']
            ax.plot(dates, Q_real, 'ro-', linewidth=1.5, markersize=4, label='Q réelle')
        
        # Ajouter les lignes de seuil
        ax.axhline(y=1200, color='red', linestyle='--', linewidth=1, alpha=0.7, label='Seuil dangereux')
        ax.axhline(y=900, color='orange', linestyle='--', linewidth=1, alpha=0.7, label='Seuil élevé')
        ax.axhline(y=600, color='green', linestyle='--', linewidth=1, alpha=0.7, label='Seuil normal')
        ax.axhline(y=400, color='yellow', linestyle='--', linewidth=1, alpha=0.7, label='Seuil modéré')
        
        # Formater l'axe des dates
        ax.xaxis.set_major_formatter(mdates.DateFormatter('%d/%m'))
        ax.xaxis.set_major_locator(mdates.DayLocator(interval=7))
        plt.xticks(rotation=45, ha='right')
        
        # Labels et titre
        ax.set_xlabel('Date', fontsize=12, fontweight='bold')
        ax.set_ylabel('Débit (m³/s)', fontsize=12, fontweight='bold')
        
        title = 'Prédictions de Débit - Saison des Pluies'
        if metadata:
            if 'k_A' in metadata and 'epsilon' in metadata:
                title += f"\nk(A) = {metadata['k_A']:.3f}, ε = {metadata['epsilon']:.2%}"
        
        ax.set_title(title, fontsize=14, fontweight='bold', pad=20)
        
        # Grille
        ax.grid(True, alpha=0.3, linestyle='--')
        
        # Légende
        ax.legend(loc='upper left', fontsize=10, framealpha=0.9)
        
        # Ajuster la mise en page
        plt.tight_layout()
        
        # Sauvegarder le graphique
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        chart_filename = f"predictions_plot_{timestamp}.png"
        chart_path = self.export_dir / chart_filename
        
        plt.savefig(chart_path, dpi=300, bbox_inches='tight')
        
        return str(chart_path)
    
    def export_coefficients(self, tableau_a: pd.DataFrame, tableau_b: pd.DataFrame, 
                          tableau_c: pd.DataFrame, filename: str = None) -> str:
        """
        Exporte les trois tableaux de coefficients en Excel
        
        Args:
            tableau_a: Tableau des coefficients annuels
            tableau_b: Tableau des coefficients mensuels
            tableau_c: Tableau des coefficients du polynôme
            filename: Nom du fichier (généré automatiquement si None)
            
        Returns:
            Chemin du fichier exporté
        """
        try:
            import openpyxl
        except ImportError:
            raise ImportError("openpyxl est requis pour l'export Excel")
        
        # Générer le nom de fichier si non fourni
        if filename is None:
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            filename = f"coefficients_{timestamp}.xlsx"
        
        if not filename.endswith('.xlsx'):
            filename += '.xlsx'
        
        filepath = self.export_dir / filename
        
        # Créer le workbook avec les trois tableaux
        with pd.ExcelWriter(filepath, engine='openpyxl') as writer:
            tableau_a.to_excel(writer, sheet_name='Tableau A - Annuel', index=False)
            tableau_b.to_excel(writer, sheet_name='Tableau B - Mensuel', index=False)
            tableau_c.to_excel(writer, sheet_name='Tableau C - Polynôme', index=False)
        
        return str(filepath)

    def export_coefficients_to_excel(self, tableau_a, tableau_b=None, 
                                    tableau_c=None, filename: str = None, metadata: dict = None) -> str:
        """
        Exporte les coefficients ou classements en Excel
        
        Args:
            tableau_a: DataFrame ou dict avec les données
            tableau_b: DataFrame optionnel (si tableau_a est DataFrame)
            tableau_c: DataFrame optionnel (si tableau_a est DataFrame)
            filename: Nom du fichier
            metadata: Métadonnées optionnelles
            
        Returns:
            Chemin du fichier exporté
        """
        try:
            import openpyxl
        except ImportError:
            raise ImportError("openpyxl est requis pour l'export Excel")
        
        # Générer le nom de fichier si non fourni
        if filename is None:
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            filename = f"export_{timestamp}.xlsx"
        
        if not filename.endswith('.xlsx'):
            filename += '.xlsx'
        
        filepath = self.export_dir / filename
        
        # Cas 1: tableau_a est un dictionnaire (classements)
        if isinstance(tableau_a, dict):
            with pd.ExcelWriter(filepath, engine='openpyxl') as writer:
                for sheet_name, df in tableau_a.items():
                    df.to_excel(writer, sheet_name=sheet_name, index=False)
                
                # Ajouter métadonnées si fournies
                if metadata:
                    meta_df = pd.DataFrame([metadata])
                    meta_df.to_excel(writer, sheet_name='Métadonnées', index=False)
        
        # Cas 2: tableau_a, b, c sont des DataFrames (coefficients)
        else:
            with pd.ExcelWriter(filepath, engine='openpyxl') as writer:
                tableau_a.to_excel(writer, sheet_name='Tableau A - Annuel', index=False)
                if tableau_b is not None:
                    tableau_b.to_excel(writer, sheet_name='Tableau B - Mensuel', index=False)
                if tableau_c is not None:
                    tableau_c.to_excel(writer, sheet_name='Tableau C - Polynôme', index=False)
        
        return str(filepath)
    
    def export_rankings_to_excel(self, annual_ranking: pd.DataFrame, 
                                monthly_ranking: pd.DataFrame, filename: str = None) -> str:
        """
        Exporte les classements annuel et mensuel en Excel
        
        Args:
            annual_ranking: Classement annuel
            monthly_ranking: Classement mensuel
            filename: Nom du fichier (généré automatiquement si None)
            
        Returns:
            Chemin du fichier exporté
        """
        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment
        except ImportError:
            raise ImportError("openpyxl est requis pour l'export Excel")
        
        # Générer le nom de fichier si non fourni
        if filename is None:
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            filename = f"classements_{timestamp}.xlsx"
        
        if not filename.endswith('.xlsx'):
            filename += '.xlsx'
        
        filepath = self.export_dir / filename
        
        # Créer le workbook avec les deux classements
        with pd.ExcelWriter(filepath, engine='openpyxl') as writer:
            annual_ranking.to_excel(writer, sheet_name='Classement Annuel', index=False)
            monthly_ranking.to_excel(writer, sheet_name='Classement Mensuel', index=False)
            
            # Formater les feuilles
            workbook = writer.book
            
            # Formater le classement annuel
            ws_annual = writer.sheets['Classement Annuel']
            header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
            header_font = Font(color="FFFFFF", bold=True)
            
            for cell in ws_annual[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal='center', vertical='center')
            
            # Formater le classement mensuel
            ws_monthly = writer.sheets['Classement Mensuel']
            for cell in ws_monthly[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal='center', vertical='center')
        
        return str(filepath)
